"""Excel-Vorlage fuer den Startbestand: erzeugen, lesen und pruefen."""
from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .kataloge import Gruppe, Merkmal
from .nummernformat import Muster, entspricht, normalisiere

SPALTEN = (
    "Inventarnummer", "Bezeichnung", "Gruppe", "Art", "Hersteller", "Typ", "Seriennummer",
    "Baujahr", "Kaufdatum", "Kaufpreis", "Lieferant", "Kostenstelle", "Menge", "Besonderheiten",
)
ARTEN = ("gross", "klein", "menge")
BLATT = "Inventar"
LISTEN = "Listen"
FEHLER_SCHLUESSEL = (
    "import.fehler.kopf_ungueltig", "import.fehler.spalte_unbekannt", "import.fehler.pflicht_fehlt",
    "import.fehler.gruppe_unbekannt", "import.fehler.art_unbekannt", "import.fehler.menge_ungueltig",
    "import.fehler.menge_nur_bei_menge", "import.fehler.kostenstelle_ungueltig",
    "import.fehler.kostenstelle_unbekannt", "import.fehler.kaufpreis_ungueltig",
    "import.fehler.kaufpreis_negativ", "import.fehler.datum_ungueltig", "import.fehler.baujahr_ungueltig",
    "import.fehler.baujahr_bereich", "import.fehler.nummer_doppelt", "import.fehler.nummer_muster",
    "import.fehler.merkmal_unbekannt_fuer_gruppe", "import.fehler.merkmal_pflicht_fehlt",
    "import.fehler.merkmal_zahl", "import.fehler.merkmal_datum", "import.fehler.merkmal_ja_nein",
    "import.fehler.merkmal_auswahl",
)
_PFLICHT = ("Inventarnummer", "Bezeichnung", "Gruppe", "Art", "Kostenstelle")


@dataclass(frozen=True)
class ImportZeile:
    zeile: int
    inventarnummer: str
    bezeichnung: str
    gruppe: str
    art: str
    hersteller: str
    typ: str
    seriennummer: str
    baujahr: int | None
    kaufdatum: date | None
    kaufpreis: Decimal | None
    lieferant: str
    kostenstelle: int
    menge: int
    besonderheiten: str
    merkmale: dict[str, str]


@dataclass(frozen=True)
class ImportFehler:
    zeile: int
    spalte: str
    text_schluessel: str
    wert: str


def _text(v: Any) -> str:
    """Macht aus einer Zelle einen getrimmten Text (ganze Zahlen ohne .0)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _ganzzahl(v: Any) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else None
    s = _text(v)
    return int(s) if re.fullmatch(r"[0-9]+", s) else None


def _zahl(v: Any) -> Decimal | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        s = str(v)
    else:
        s = _text(v).replace(" ", "")
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
        else:
            s = s.replace(",", ".")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


def _datum(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    for muster in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, muster).date()
        except ValueError:
            continue
    return None


def _merkmal_wert(m: Merkmal, v: Any) -> tuple[str | None, str]:
    """Prueft einen Merkmalswert; liefert (Wert, Fehlerschluessel) - genau eines ist leer."""
    if m.typ == "zahl":
        z = _zahl(v)
        return (None, "import.fehler.merkmal_zahl") if z is None else (format(z, "f"), "")
    if m.typ == "datum":
        d = _datum(v)
        return (None, "import.fehler.merkmal_datum") if d is None else (d.isoformat(), "")
    if m.typ == "ja_nein":
        s = "ja" if v is True else "nein" if v is False else _text(v).lower()
        if s in ("ja", "j", "1", "true"):
            return "ja", ""
        if s in ("nein", "n", "0", "false"):
            return "nein", ""
        return None, "import.fehler.merkmal_ja_nein"
    s = _text(v)
    if m.typ == "auswahl" and s not in m.auswahl:
        return None, "import.fehler.merkmal_auswahl"
    return s, ""


def _beispielwert(m: Merkmal) -> Any:
    return {"zahl": 1, "datum": "01.01.2026", "ja_nein": "ja", "auswahl": (m.auswahl or ("",))[0]}.get(m.typ, "Beispiel")


def erzeuge_vorlage(pfad: Path, gruppen: Iterable[Gruppe], merkmale: Iterable[Merkmal]) -> None:
    """Schreibt die Vorlage: Kopfzeile, zwei Beispielzeilen, Dropdowns, Merkmalspalten mit Kommentar."""
    gruppen = list(gruppen)
    merkmale = list(merkmale)
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    spalten_m: dict[str, list[Merkmal]] = {}
    for m in merkmale:
        spalten_m.setdefault(m.schluessel, []).append(m)
    kopf = list(SPALTEN) + [f"m:{s}" for s in spalten_m]
    ws.append(kopf)
    for i, name in enumerate(kopf, start=1):
        zelle = ws.cell(row=1, column=i)
        zelle.font = Font(bold=True)
        ws.column_dimensions[zelle.column_letter].width = max(14, len(name) + 2)
        if name.startswith("m:"):
            zelle.comment = Comment("\n".join(
                f"gruppe={m.gruppe}; typ={m.typ}; einheit={m.einheit}; pflicht={int(m.pflicht)}; "
                f"auswahl={'|'.join(m.auswahl)}" for m in spalten_m[name[2:]]), "inventar_rein")
    ws.freeze_panes = "A2"
    for nr, (g, art, menge) in enumerate(
        [(gruppen[0], "gross", 1), (gruppen[1] if len(gruppen) > 1 else gruppen[0], "menge", 40)], start=2
    ):
        zeile: dict[str, Any] = {
            "Inventarnummer": f"{g.kuerzel}-00001", "Bezeichnung": f"Beispiel {g.bezeichnung}",
            "Gruppe": g.schluessel, "Art": art, "Hersteller": "Beispiel-Hersteller", "Typ": "BSP-100",
            "Seriennummer": "TEST-000001", "Baujahr": 2020, "Kaufdatum": "15.03.2020", "Kaufpreis": 1000,
            "Lieferant": "Beispiel-Lieferant", "Kostenstelle": 1000, "Menge": menge, "Besonderheiten": "",
        }
        for m in merkmale:
            if m.gruppe == g.schluessel and m.pflicht:
                zeile[f"m:{m.schluessel}"] = _beispielwert(m)
        ws.append([zeile.get(name, None) for name in kopf])
    listen = wb.create_sheet(LISTEN)
    listen.append(["Gruppe", "Art"])
    for i in range(max(len(gruppen), len(ARTEN))):
        listen.append([gruppen[i].schluessel if i < len(gruppen) else None, ARTEN[i] if i < len(ARTEN) else None])
    listen.sheet_state = "hidden"
    for spalte, bis in (("C", len(gruppen) + 1), ("D", len(ARTEN) + 1)):
        dv = DataValidation(type="list", formula1=f"={LISTEN}!${'A' if spalte == 'C' else 'B'}$2:${'A' if spalte == 'C' else 'B'}${bis}",
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{spalte}2:{spalte}2000")
    wb.save(pfad)


def lies(
    pfad: Path, gruppen: Iterable[Gruppe], merkmale: Iterable[Merkmal], kostenstellen: Iterable[int],
    muster: Muster | None, heute: date | None = None,
) -> tuple[tuple[ImportZeile, ...], tuple[ImportFehler, ...]]:
    """Liest und prueft die Datei; fehlerhafte Zeilen kommen nur in die Fehlerliste."""
    gruppen_schluessel = {g.schluessel for g in gruppen}
    je_gruppe: dict[str, dict[str, Merkmal]] = {}
    for m in merkmale:
        je_gruppe.setdefault(m.gruppe, {})[m.schluessel] = m
    erlaubte_ks = frozenset(kostenstellen)
    jahr_max = (heute or date.today()).year
    wb = load_workbook(pfad, data_only=True)
    ws = wb[BLATT] if BLATT in wb.sheetnames else wb.worksheets[0]
    zeilen = list(ws.iter_rows(values_only=True))
    kopf = [_text(c) for c in (zeilen[0] if zeilen else ())]
    while kopf and not kopf[-1]:
        kopf.pop()
    fehler: list[ImportFehler] = []
    if tuple(kopf[: len(SPALTEN)]) != SPALTEN:
        return (), (ImportFehler(1, "", "import.fehler.kopf_ungueltig", ";".join(kopf)),)
    fehler += [ImportFehler(1, name, "import.fehler.spalte_unbekannt", name)
               for name in kopf[len(SPALTEN):] if not name.startswith("m:") or len(name) == 2]
    ok: dict[int, ImportZeile] = {}
    for nr, roh in enumerate(zeilen[1:], start=2):
        werte = {name: (roh[i] if i < len(roh) else None) for i, name in enumerate(kopf)}
        if all(_text(v) == "" for v in werte.values()):
            continue
        zf: list[ImportFehler] = []

        def fehl(spalte: str, schluessel: str, v: Any = None) -> None:
            zf.append(ImportFehler(nr, spalte, schluessel, _text(werte.get(spalte) if v is None else v)))

        for spalte in _PFLICHT:
            if _text(werte[spalte]) == "":
                fehl(spalte, "import.fehler.pflicht_fehlt")
        nummer = normalisiere(_text(werte["Inventarnummer"]))
        if nummer and muster is not None and not entspricht(muster, nummer):
            fehl("Inventarnummer", "import.fehler.nummer_muster", nummer)
        gruppe, art = _text(werte["Gruppe"]), _text(werte["Art"])
        if gruppe and gruppe not in gruppen_schluessel:
            fehl("Gruppe", "import.fehler.gruppe_unbekannt")
        if art and art not in ARTEN:
            fehl("Art", "import.fehler.art_unbekannt")
        ks = _ganzzahl(werte["Kostenstelle"])
        if _text(werte["Kostenstelle"]) != "":
            if ks is None:
                fehl("Kostenstelle", "import.fehler.kostenstelle_ungueltig")
            elif ks not in erlaubte_ks:
                fehl("Kostenstelle", "import.fehler.kostenstelle_unbekannt")
        menge = 1
        if art == "menge":
            menge_roh = _ganzzahl(werte["Menge"])
            if menge_roh is None or menge_roh < 1:
                fehl("Menge", "import.fehler.menge_ungueltig")
            else:
                menge = menge_roh
        elif _text(werte["Menge"]) != "":
            menge_roh = _ganzzahl(werte["Menge"])
            if menge_roh is None:
                fehl("Menge", "import.fehler.menge_ungueltig")
            elif menge_roh != 1:
                fehl("Menge", "import.fehler.menge_nur_bei_menge")
        baujahr = None
        if _text(werte["Baujahr"]) != "":
            baujahr = _ganzzahl(werte["Baujahr"])
            if baujahr is None:
                fehl("Baujahr", "import.fehler.baujahr_ungueltig")
            elif not 1950 <= baujahr <= jahr_max:
                fehl("Baujahr", "import.fehler.baujahr_bereich")
        kaufdatum = None
        if _text(werte["Kaufdatum"]) != "":
            kaufdatum = _datum(werte["Kaufdatum"])
            if kaufdatum is None:
                fehl("Kaufdatum", "import.fehler.datum_ungueltig")
        kaufpreis = None
        if _text(werte["Kaufpreis"]) != "":
            kaufpreis = _zahl(werte["Kaufpreis"])
            if kaufpreis is None:
                fehl("Kaufpreis", "import.fehler.kaufpreis_ungueltig")
            elif kaufpreis < 0:
                fehl("Kaufpreis", "import.fehler.kaufpreis_negativ")
        werte_m: dict[str, str] = {}
        erlaubt = je_gruppe.get(gruppe, {})
        for name in kopf[len(SPALTEN):]:
            if not name.startswith("m:") or _text(werte[name]) == "":
                continue
            m = erlaubt.get(name[2:])
            if m is None:
                fehl(name, "import.fehler.merkmal_unbekannt_fuer_gruppe")
                continue
            wert, schluessel = _merkmal_wert(m, werte[name])
            if wert is None:
                fehl(name, schluessel)
            else:
                werte_m[m.schluessel] = wert
        for m in erlaubt.values():
            if m.pflicht and m.schluessel not in werte_m and not any(
                f.spalte == f"m:{m.schluessel}" for f in zf
            ):
                zf.append(ImportFehler(nr, f"m:{m.schluessel}", "import.fehler.merkmal_pflicht_fehlt", ""))
        fehler += zf
        if not zf:
            ok[nr] = ImportZeile(
                nr, nummer, _text(werte["Bezeichnung"]), gruppe, art, _text(werte["Hersteller"]),
                _text(werte["Typ"]), _text(werte["Seriennummer"]), baujahr, kaufdatum, kaufpreis,
                _text(werte["Lieferant"]), ks or 0, menge, _text(werte["Besonderheiten"]), werte_m)
    nummern: dict[str, list[int]] = {}
    for z in _nummern(ws, ok):
        nummern.setdefault(z[1], []).append(z[0])
    for nummer, rows in nummern.items():
        if len(rows) > 1:
            for r in rows:
                fehler.append(ImportFehler(r, "Inventarnummer", "import.fehler.nummer_doppelt", nummer))
                ok.pop(r, None)
    fehler.sort(key=lambda f: (f.zeile, f.spalte, f.text_schluessel))
    return tuple(ok[k] for k in sorted(ok)), tuple(fehler)


def _nummern(ws: Any, ok: dict[int, ImportZeile]) -> list[tuple[int, str]]:
    """Alle normalisierten Inventarnummern der Datei (auch von fehlerhaften Zeilen) mit Zeilennummer."""
    ergebnis = []
    for nr, roh in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nummer = normalisiere(_text(roh[0] if roh else None))
        if nummer:
            ergebnis.append((nr, nummer))
    return ergebnis
