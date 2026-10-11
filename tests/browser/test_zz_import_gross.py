"""Browserfall 4 (INTEGRATION_VSC.md Abschnitt 4): Import mit 2.500 Stücken, danach öffnet `/inventar` in unter einer Sekunde.

Der Fall legt 2.500 Stücke an und steht darum **zuletzt** (Dateiname mit `zz`); eine zweite Fahrt braucht eine frische Box (die Nummern gibt es dann schon).
"""

from __future__ import annotations

import time

from openpyxl import load_workbook
from playwright.sync_api import Page, expect

ANZAHL = 2500


def test_F04_import_2500_stuecke_und_liste_unter_einer_sekunde(verwalter: Page, tmp_path) -> None:
    verwalter.goto("/inventar/verwaltung/import")
    with verwalter.expect_download() as geladen:
        verwalter.click("a[href='/inventar/verwaltung/import/vorlage']")
    vorlage = tmp_path / "vorlage.xlsx"
    geladen.value.save_as(vorlage)
    mappe = load_workbook(vorlage)
    blatt = mappe["Inventar"]
    kopf = [z.value for z in blatt[1]]
    blatt.delete_rows(2, blatt.max_row)  # die Beispielzeilen weg
    for i in range(1, ANZAHL + 1):
        zeile = {"Inventarnummer": f"KG-{90000 + i:05d}", "Bezeichnung": f"Importstück {i}", "Gruppe": "kleingeraet", "Art": "klein",
                 "Kostenstelle": 70101}
        blatt.append([zeile.get(name) for name in kopf])
    datei = tmp_path / "import_2500.xlsx"
    mappe.save(datei)
    verwalter.set_input_files("#import-datei", str(datei))
    verwalter.click("main form[enctype='multipart/form-data'] button.haupt", no_wait_after=True)
    expect(verwalter.locator("main")).to_contain_text(f"{ANZAHL}", timeout=120_000)
    expect(verwalter.locator("table.liste .schlimm")).to_have_count(0)
    verwalter.click("main form:has(input[value=einspielen]) button.haupt", no_wait_after=True)
    expect(verwalter.locator("p.gut[role=status]")).to_contain_text(str(ANZAHL), timeout=600_000)
    dauern = []
    for weg in ("/inventar", "/inventar?ansicht=liste"):
        anfang = time.perf_counter()
        verwalter.goto(weg, wait_until="load")
        dauern.append(time.perf_counter() - anfang)
    assert max(dauern) < 1.0, dauern
